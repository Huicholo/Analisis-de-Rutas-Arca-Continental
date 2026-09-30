"""Motor de cálculo del análisis de Madurez × Atractividad (réplica en Python de la hoja
'Analisis recomendación Rutas v2' del libro de análisis).

El libro de Excel sólo aporta datos y parámetros:
- 'Base de datos Rutas'   → datos por ruta y los bloques "Atractividad Caso NN" (casos de atractividad).
- 'Hoja de apoyo v2'      → bloques de estándares por Región (Conservador / Medio / Ambicioso…).
- 'Rentabilidad datos'    → venta y costo de servir por CeDi.
- 'Análisis de KPIs'      → supuestos para calcular los pesos del Índice de Ejecución.
- 'Analisis recomendación Rutas v2' → sólo sus celdas de parámetros (filas 3 y 4).

Todas las fórmulas (índices, puntajes, pesos, saturación, cortes y P) se calculan aquí.
Cada bloque indica la columna del Excel que replica para poder auditarlo.
"""
from __future__ import annotations

import math
import re
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

# --------------------------------------------------------------------------- nombres de hojas
HOJA_RUTAS = "Base de datos Rutas"
HOJA_APOYO = "Hoja de apoyo v2"
HOJA_RENTABILIDAD = "Rentabilidad datos"
HOJA_KPIS = "Análisis de KPIs"
HOJA_ANALISIS = "Analisis recomendación Rutas v2"

FILA_ENCABEZADO_RUTAS = 5      # fila de encabezados en 'Base de datos Rutas'
FILA_BLOQUES_RUTAS = 3         # fila con "Atractividad Caso NN"

# Parámetros editables en la hoja de análisis (celda → nombre). Si una celda está vacía se usa el default.
PARAMETROS_ANALISIS = {
    "cortes_cumplimiento": (["U3", "V3", "W3", "X3"], [0.5, 0.75, 1.0, 1.25]),
    "min_cobertura_madurez": ("AI4", 0.3),
    "corte_madura": ("AJ4", 65.0),
    "corte_inmadura": ("AK4", 45.0),
    "confiabilidad_alta": ("AL3", 0.9),
    "confiabilidad_media": ("AL4", 0.6),
    "corte_riesgo_cds": ("AO4", 0.6),
    "holgura_max": ("BC4", 0.0),
    "pesos_atractivo": (["BK3", "BL3", "BM3", "BN3"], [0.25, 0.25, 0.25, 0.25]),
    "min_cobertura_zona": ("BO4", 0.3),
    "corte_atractivo_alto": ("BP4", 50.0),
    "corte_atractivo_bajo": ("BQ4", 25.0),
    "peso_ejecucion": ("BS3", 0.75),
    "peso_saturacion": ("BU3", 0.25),
}
# Supuestos de 'Análisis de KPIs' para los pesos del Índice de Ejecución
PARAMETROS_KPIS = {
    "peso_impacto": ("B30", 0.4),
    "peso_no_redundancia": ("B31", 0.2),
    "peso_estrategico": ("B32", 0.4),
    "factor_op": ("B34", 3.0),
    "factor_gfn": ("C34", 2.0),
    "relevancia": (["C38", "C39", "C40", "C41", "C42", "C43", "C44"], [5, 4, 4, 4, 3, 3, 3]),
    "buckets": (["B8", "B9", "B10", "B11", "B12", "B13", "B14"],
                ["Ruta", "Ruta", "Ingreso", "Ingreso", "Volumen", "Volumen", "Consumo"]),
}

# Los 7 KPI del Índice de Ejecución (orden de la hoja de análisis, columnas G:M)
KPIS = ["Clientes/día", "Efectividad de compra", "Ingreso OP/día", "Ingreso GFN/día",
        "#CU GFN/día", "#CU OP/día", "Ticket x cliente"]
KPIS_EST = ["Est. Clientes/día", "Est. Efectividad", "Est. Ingreso OP/día", "Est. Ingreso GFN/día",
            "Est. #CU GFN/día", "Est. #CU OP/día", "Est. Ticket"]
KPIS_CUMPL = ["Cumpl. Clientes/día", "Cumpl. Efectividad", "Cumpl. Ingreso OP/día", "Cumpl. Ingreso GFN/día",
              "Cumpl. #CU GFN/día", "Cumpl. #CU OP/día", "Cumpl. Ticket"]
KPIS_SC = ["Sc. Clientes/día", "Sc. Efectividad", "Sc. Ingreso OP/día", "Sc. Ingreso GFN/día",
           "Sc. #CU GFN/día", "Sc. #CU OP/día", "Sc. Ticket"]
ZONA = ["% Hogares totales ABC+", "Hogares actual + potencial (%)", "Potencial Ingreso / hogar", "Potencial CU / hogar"]
ZONA_SC = ["Sc. % Hogares ABC+", "Sc. Hogares act.+pot.", "Sc. Ingreso / hogar", "Sc. CU / hogar"]

# Columnas de 'Base de datos Rutas' que usa el modelo (por encabezado de la fila 5)
COLS_RUTAS = {
    "Región": "region", "Territorio EERR": "territorio", "CeDi": "cedi", "Ruta": "ruta",
    "Estatus de ruta": "estatus", "Elegible para estándar": "elegible",
    "Rutas Hogar del CeDi": "rutas_cedi",
    "Cajas Unidad Garrafón": "cu_gfn", "Cajas Unidad Otros Productos": "cu_op",
    "Volumen Total en Cajas Unidad [CU]": "cu_total",
    "Venta GFN [MXN]": "venta_gfn", "Venta OP [MXN]": "venta_op",
    "Ticket $ / CC": "ticket",
    "Clientes atiende/día (real)": "clientes_dia", "Efectividad de compra (real)": "efectividad",
    "Volumen CU/día (real)": "cu_dia",
    "Hogares": "hogares",
    "Clientes actuales (Total)": "clientes_act", "Clientes potencial": "clientes_pot",
    "Horas actuales": "horas_act", "Jornada completa": "jornada",
    "Capacidad instalada": "cap_inst", "Capacidad Utilizada": "cap_util",
    # Sólo para las recomendaciones del árbol de decisión (no entran al cálculo de la P)
    "Hogares por manzana": "hog_manzana", "% A/B + C+": "pct_abc_nse",
    "Capacidad adicional (clientes)": "cap_adicional", "% Venta OP (ingreso)": "pct_venta_op",
    "% Venta Web.": "pct_venta_web",
}
# Dentro de cada bloque "Atractividad Caso NN": posición (0-based) de cada insumo capturado
OFFSETS_CASO = {"hog_abc": 0, "hog_pot": 2, "pot_ingreso": 5, "pot_cu": 7}

# Columnas de estándares en 'Hoja de apoyo v2' (encabezado del bloque) → columna del modelo
COLS_ESTANDAR = {
    "Clientes/día": "Est. Clientes/día", "Efectividad de compra": "Est. Efectividad",
    "Ingreso OP/día": "Est. Ingreso OP/día", "Ingreso GFN/día": "Est. Ingreso GFN/día",
    "#CU GFN/día": "Est. #CU GFN/día", "#CU OP/día": "Est. #CU OP/día", "Ticket x cliente": "Est. Ticket",
}
# Estándares que Excel calcula como p75 de la propia base (no dependen del escenario)
ESTANDAR_P75 = {"Est. Clientes/día": "Clientes/día", "Est. Efectividad": "Efectividad de compra",
                "Est. Ticket": "Ticket x cliente"}

SIN_DATOS = "Sin datos"
COL_HOG_ABC_TOTAL = "Hogares A/B + C+ (total)"
COL_HOG_ABC_POT = "Hogares A/B + C+ potenciales a capturar"
COL_HOG_ABC_ACT = "Hogares A/B + C+ actuales"
COL_WEB = "% Venta Web"
COL_ABC_CAPTURABLE = "% ABC+ capturable"          # hogares A/B + C+ por capturar ÷ hogares A/B + C+ totales
COL_OP_VOLUMEN = "% OP (volumen)"                  # cajas unidad de otros productos ÷ cajas unidad totales


# --------------------------------------------------------------------------- utilidades
def _num(v) -> float:
    """ISNUMBER de Excel: sólo números (no texto, no booleanos, no errores)."""
    if isinstance(v, bool) or v is None:
        return math.nan
    if isinstance(v, (int, float, np.integer, np.floating)):
        return float(v)
    return math.nan


def _n(v) -> float:
    """N() de Excel / celda vacía en aritmética: número o 0."""
    x = _num(v)
    return 0.0 if math.isnan(x) else x


def _redondear(x: float, dec: int = 1) -> float:
    """ROUND de Excel (mitades se alejan de cero), no el redondeo bancario de Python."""
    if x is None or (isinstance(x, float) and math.isnan(x)):
        return math.nan
    f = 10 ** dec
    return math.copysign(math.floor(abs(x) * f + 0.5) / f, x)


def _p75(valores: pd.Series) -> float:
    """AGGREGATE(16,6,…,0.75) = PERCENTILE.INC ignorando no numéricos."""
    v = pd.to_numeric(valores, errors="coerce").dropna()
    return float(np.percentile(v, 75)) if len(v) else math.nan


def _div(a, b) -> float:
    try:
        a, b = float(a), float(b)
    except (TypeError, ValueError):
        return math.nan
    if math.isnan(a) or math.isnan(b) or b == 0:
        return math.nan
    return a / b


def _celda(ws, ref, default):
    if isinstance(ref, list):
        vals = [_num(ws[r].value) if not isinstance(d, str) else ws[r].value for r, d in zip(ref, default)]
        return [d if (v is None or (isinstance(v, float) and math.isnan(v))) else v for v, d in zip(vals, default)]
    v = ws[ref].value
    x = _num(v)
    return default if math.isnan(x) else x


# --------------------------------------------------------------------------- lectura del libro
@dataclass
class Insumos:
    rutas: pd.DataFrame                    # una fila por renglón de 'Base de datos Rutas'
    casos_atractividad: dict[str, pd.DataFrame]   # nombre → columnas hog_abc, hog_pot, pot_ingreso, pot_cu
    escenarios: dict[str, pd.DataFrame]    # nombre → estándares por Región
    rentabilidad: pd.DataFrame             # por CeDi: venta_ruta, pct_cds
    params: dict = field(default_factory=dict)
    params_kpi: dict = field(default_factory=dict)
    segmentos: "Segmentos | None" = None

    def caso_disponible(self, caso: str) -> bool:
        d = self.casos_atractividad[caso]
        return bool(d.apply(pd.to_numeric, errors="coerce").notna().any().any())


def _leer_rutas(ws) -> tuple[pd.DataFrame, dict[str, pd.DataFrame]]:
    filas = list(ws.iter_rows(min_row=FILA_BLOQUES_RUTAS, values_only=True))
    bloques = filas[0]
    enc = filas[FILA_ENCABEZADO_RUTAS - FILA_BLOQUES_RUTAS]
    datos = filas[FILA_ENCABEZADO_RUTAS - FILA_BLOQUES_RUTAS + 1:]

    idx = {}
    for i, h in enumerate(enc):
        if isinstance(h, str) and h.strip() in COLS_RUTAS and COLS_RUTAS[h.strip()] not in idx:
            idx[COLS_RUTAS[h.strip()]] = i
    faltan = [k for k, v in COLS_RUTAS.items() if v not in idx]
    if faltan:
        raise ValueError(f"En '{HOJA_RUTAS}' faltan columnas: {', '.join(faltan)}")

    datos = [f for f in datos if f and f[idx["region"]] not in (None, "")]
    rutas = pd.DataFrame({k: [f[i] for f in datos] for k, i in idx.items()})

    casos = {}
    for i, b in enumerate(bloques):
        m = re.match(r"\s*Atractividad\s+Caso\s+(.+)", str(b or ""), flags=re.I)
        if m:
            nombre = f"Caso {m.group(1).strip()}"
            casos[nombre] = pd.DataFrame({k: [f[i + off] if i + off < len(f) else None for f in datos]
                                          for k, off in OFFSETS_CASO.items()})
    if not casos:
        raise ValueError(f"En '{HOJA_RUTAS}' no hay bloques 'Atractividad Caso …' en la fila {FILA_BLOQUES_RUTAS}")
    return rutas, casos


def _leer_escenarios(ws) -> dict[str, pd.DataFrame]:
    """Bloques de 'Hoja de apoyo v2': una fila con el nombre del escenario, luego encabezados
    ('Región', 'Clientes/día', …) y una fila por Región hasta una fila vacía."""
    filas = list(ws.iter_rows(values_only=True))
    escenarios = {}
    for r in range(len(filas) - 1):
        nombre, siguiente = filas[r][0], filas[r + 1]
        if (isinstance(nombre, str) and nombre.strip() and siguiente and siguiente[0] == "Región"
                and "Clientes/día" in siguiente and "Ingreso OP/día" in siguiente):
            enc = siguiente
            regs = []
            for f in filas[r + 2:]:
                if not f or f[0] in (None, ""):
                    break
                regs.append({"Región": f[0], **{COLS_ESTANDAR[h]: f[j] for j, h in enumerate(enc)
                                                 if h in COLS_ESTANDAR}})
            escenarios[nombre.strip()] = pd.DataFrame(regs)
    if not escenarios:
        raise ValueError(f"En '{HOJA_APOYO}' no se encontraron bloques de estándares por Región")
    return escenarios


def _leer_rentabilidad(ws) -> pd.DataFrame:
    filas = list(ws.iter_rows(values_only=True))
    enc = [str(h).strip() if h is not None else "" for h in filas[0]]
    pos = {n: enc.index(n) for n in ["CeDi", "Rutas Hogar", "Venta Neta", "Costo-De-Servir x Cedis"]}
    reg = []
    for f in filas[1:]:
        cedi = f[pos["CeDi"]]
        if cedi in (None, ""):
            continue
        d, l, n = f[pos["Rutas Hogar"]], f[pos["Venta Neta"]], f[pos["Costo-De-Servir x Cedis"]]
        reg.append({
            "cedi": cedi,
            "venta_ruta": _div(_num(l), _num(d)),                                  # M = L/D
            "pct_cds": math.nan if n in (None, "") else _div(_num(n), _num(l)),    # P = N/L
        })
    # MATCH devuelve la primera coincidencia
    return pd.DataFrame(reg, columns=["cedi", "venta_ruta", "pct_cds"]).drop_duplicates("cedi", keep="first")


@dataclass
class Segmentos:
    """Variables extra de la hoja de segmentos, listas para cruzar con las rutas."""
    hoja: str
    llave: str                   # "Ruta" o "CeDi": columna del modelo con la que se cruza
    datos: pd.DataFrame          # columna llave + una columna numérica por variable
    variables: list[str]


def _normaliza_texto(t) -> str:
    t = unicodedata.normalize("NFKD", str(t or "")).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", t).strip().lower()


def _leer_segmentos(wb) -> Segmentos | None:
    """Hoja cuyo nombre contiene 'segmento'. Encabezados: la primera fila (de las 15 primeras) con
    una celda 'Ruta' o 'CeDi'. Cada columna mayoritariamente numérica se vuelve una variable."""
    hoja = next((h for h in wb.sheetnames if "segmento" in _normaliza_texto(h)), None)
    if hoja is None:
        return None
    filas = list(wb[hoja].iter_rows(values_only=True))
    for r, fila in enumerate(filas[:15]):
        normal = [_normaliza_texto(c) for c in fila]
        llave = "Ruta" if "ruta" in normal else ("CeDi" if "cedi" in normal else None)
        if llave:
            break
    else:
        raise ValueError(f"En '{hoja}' no encontré una fila de encabezados con 'Ruta' o 'CeDi' (primeras 15 filas)")
    pos_llave = normal.index(llave.lower())
    nombres, vistos = [], set()
    for i, h in enumerate(fila):
        n = str(h).strip() if h not in (None, "") else f"Columna {i + 1}"
        while n in vistos:
            n += " (2)"
        vistos.add(n)
        nombres.append(n)
    datos = [f for f in filas[r + 1:] if f and pos_llave < len(f) and f[pos_llave] not in (None, "")]
    df = pd.DataFrame([list(f) + [None] * (len(nombres) - len(f)) for f in datos], columns=nombres[:max(len(nombres), 1)])
    ident = {"ruta", "cedi", "region", "territorio", "territorio eerr", "ruta id", "estatus", "estatus de ruta"}
    variables = []
    for col in df.columns:
        if col == nombres[pos_llave] or _normaliza_texto(col) in ident:
            continue
        no_nulos = df[col].dropna()
        numeros = no_nulos.map(lambda v: not math.isnan(_num(v)))
        if len(no_nulos) and numeros.mean() >= 0.5:
            df[col] = df[col].map(_num)
            variables.append(col)
    clave = df[nombres[pos_llave]]
    if llave == "Ruta":
        clave = clave.map(lambda v: str(int(v)) if not math.isnan(_num(v)) and float(v).is_integer() else str(v).strip())
    else:
        clave = clave.map(lambda v: str(v).strip())
    salida = pd.DataFrame({llave: clave}).join(df[variables])
    # Si la llave se repite, se usa la primera fila (como MATCH)
    salida = salida.drop_duplicates(llave, keep="first").reset_index(drop=True)
    return Segmentos(hoja=hoja, llave=llave, datos=salida, variables=variables)


def leer_libro(ruta: str | Path) -> Insumos:
    wb = openpyxl.load_workbook(ruta, data_only=True, read_only=False)
    faltan = [h for h in (HOJA_RUTAS, HOJA_APOYO, HOJA_RENTABILIDAD, HOJA_KPIS, HOJA_ANALISIS)
              if h not in wb.sheetnames]
    if faltan:
        raise ValueError(f"El libro no tiene las hojas: {', '.join(faltan)}")
    rutas, casos = _leer_rutas(wb[HOJA_RUTAS])
    ws_a, ws_k = wb[HOJA_ANALISIS], wb[HOJA_KPIS]
    return Insumos(
        rutas=rutas,
        casos_atractividad=casos,
        escenarios=_leer_escenarios(wb[HOJA_APOYO]),
        rentabilidad=_leer_rentabilidad(wb[HOJA_RENTABILIDAD]),
        params={k: _celda(ws_a, ref, d) for k, (ref, d) in PARAMETROS_ANALISIS.items()},
        params_kpi={k: _celda(ws_k, ref, d) for k, (ref, d) in PARAMETROS_KPIS.items()},
        segmentos=_leer_segmentos(wb),
    )


# --------------------------------------------------------------------------- cálculo
def _pesos_ejecucion(df: pd.DataFrame, pk: dict) -> list[float]:
    """Réplica de 'Análisis de KPIs' secciones A–D → pesos AB3:AH3."""
    base = df[(df["Estatus"] == "Activa") & (df["Elegible p/ estándar"] == "Sí")]
    kp = base[KPIS].apply(pd.to_numeric, errors="coerce")
    margen = pd.to_numeric(base["Margen estimado"], errors="coerce")
    venta = pd.to_numeric(base["Venta Neta de la ruta"], errors="coerce")

    impacto = [np.mean([abs(kp[k].corr(margen)), abs(kp[k].corr(venta))]) for k in KPIS]      # D
    corr = kp.corr()                                                                           # matriz B
    redund = [(corr.loc[k].abs().sum() - 1) / (len(KPIS) - 1) for k in KPIS]                   # I19:I25
    no_red = [1 - r for r in redund]                                                           # E
    rel = [float(x) for x in pk["relevancia"]]

    f_ = [d / sum(impacto) for d in impacto]
    g_ = [e / sum(no_red) for e in no_red]
    h_ = [c / sum(rel) for c in rel]
    base_w = [f * pk["peso_impacto"] + g * pk["peso_no_redundancia"] + h * pk["peso_estrategico"]
              for f, g, h in zip(f_, g_, h_)]                                                  # I
    factor = []
    for k in KPIS:                                                                             # J (SEARCH no distingue mayúsculas)
        ku = k.upper()
        factor.append(pk["factor_op"] if "OP" in ku else (pk["factor_gfn"] if "GFN" in ku else None))
    buckets = pk["buckets"]
    pesos = []
    for i in range(len(KPIS)):                                                                 # K
        if factor[i] is None:
            pesos.append(base_w[i])
        else:
            mismos = [j for j in range(len(KPIS)) if buckets[j] == buckets[i]]
            suma_i = sum(base_w[j] for j in mismos)
            suma_f = sum(factor[j] for j in mismos if factor[j] is not None)
            pesos.append(suma_i * factor[i] / suma_f)
    return pesos


def _puntaje_quintil_cedi(df: pd.DataFrame, col: str) -> pd.Series:
    """BK:BN → MIN(5, 1+INT(5 * #(mismo CeDi y valor < x) / #(mismo CeDi con número)))."""
    out = pd.Series(0.0, index=df.index)
    vals = pd.to_numeric(df[col], errors="coerce")
    for _, idx in df.groupby("CeDi", sort=False).groups.items():
        v = vals.loc[idx]
        n = v.notna().sum()
        for i in idx:
            x = vals.at[i]
            if math.isnan(x):
                continue
            menores = (v < x).sum()
            out.at[i] = min(5, 1 + math.floor(5 * menores / n))
    return out


def calcular(ins: Insumos, escenario: str, caso: str) -> tuple[pd.DataFrame, dict]:
    """Devuelve (tabla con las columnas de la hoja de análisis, resumen de cortes y pesos)."""
    if escenario not in ins.escenarios:
        raise ValueError(f"Escenario '{escenario}' no existe. Opciones: {list(ins.escenarios)}")
    if caso not in ins.casos_atractividad:
        raise ValueError(f"Caso '{caso}' no existe. Opciones: {list(ins.casos_atractividad)}")
    p = ins.params
    r = ins.rutas
    n = len(r)
    df = pd.DataFrame(index=range(n))

    # 1. Identificación (A:F)
    df["Región"] = r["region"]
    df["Territorio EERR"] = r["territorio"]
    df["CeDi"] = r["cedi"]
    # Ruta como texto ("105300"); sin código válido → "Sin código" (columna D de la hoja de análisis)
    df["Ruta"] = ["Sin código" if math.isnan(_num(x)) else
                  (str(int(x)) if float(x).is_integer() else str(x)) for x in r["ruta"]]
    df["Estatus"] = r["estatus"]
    df["Elegible p/ estándar"] = r["elegible"]

    # 2. Indicador actual (G:M). Ingresos y CU por día = total / (CU total / CU por día)
    df["Clientes/día"] = r["clientes_dia"].map(_num)
    df["Efectividad de compra"] = r["efectividad"].map(_num)
    dias = [(_div(_num(t), _n(d)) if _n(d) != 0 else math.nan) for t, d in zip(r["cu_total"], r["cu_dia"])]
    for col, src in [("Ingreso OP/día", "venta_op"), ("Ingreso GFN/día", "venta_gfn"),
                     ("#CU GFN/día", "cu_gfn"), ("#CU OP/día", "cu_op")]:
        df[col] = [_div(_num(v), d) for v, d in zip(r[src], dias)]
    df["Ticket x cliente"] = r["ticket"].map(_num)

    # 7. Rentabilidad (AM:AS) — se necesita antes para el p75 de margen y los pesos
    rent = ins.rentabilidad.set_index("cedi")
    venta_ruta = r["cedi"].map(rent["venta_ruta"])            # Base W
    pct_cds = r["cedi"].map(rent["pct_cds"])                  # Base X
    df["Venta Neta de la ruta"] = venta_ruta.astype(float)
    suma_cedi = venta_ruta.groupby(r["cedi"]).transform(lambda s: s.fillna(0).sum())
    df["Ingreso prom. x ruta del CeDi"] = [_div(s, _num(rc)) for s, rc in zip(suma_cedi, r["rutas_cedi"])]
    df["%CDS asignado a la ruta"] = [_div(x * an, am) if not math.isnan(x) else math.nan
                                     for x, an, am in zip(pct_cds.astype(float), df["Ingreso prom. x ruta del CeDi"],
                                                          df["Venta Neta de la ruta"])]
    df["Margen estimado"] = 1 - df["%CDS asignado a la ruta"]

    # 3. Estándar por Región (N:T): p75 propios + bloque del escenario
    eleg = df["Elegible p/ estándar"] == "Sí"
    est = ins.escenarios[escenario].set_index("Región")
    for col_est in KPIS_EST:
        if col_est in ESTANDAR_P75:
            fuente = ESTANDAR_P75[col_est]
            mapa = {reg: _p75(df.loc[eleg & (df["Región"] == reg), fuente]) for reg in df["Región"].unique()}
            df[col_est] = df["Región"].map(mapa)
        else:
            df[col_est] = df["Región"].map(est[col_est].map(_num)) if col_est in est else math.nan
    margen_bench = {reg: _p75(df.loc[eleg & (df["Región"] == reg), "Margen estimado"]) for reg in df["Región"].unique()}
    df["Margen bench = p75 de su Región"] = df["Región"].map(margen_bench)
    df["Margen vs. bench (pp)"] = df["Margen estimado"] - df["Margen bench = p75 de su Región"]
    df["¿En riesgo?"] = ["n/d" if math.isnan(x) else ("Sí" if x >= p["corte_riesgo_cds"] else "No")
                         for x in df["%CDS asignado a la ruta"]]

    # 4. Cumplimiento (U:AA)
    for k, e, c in zip(KPIS, KPIS_EST, KPIS_CUMPL):
        df[c] = [_div(a, b) for a, b in zip(df[k], df[e])]

    # 5. Puntaje 1-5 (AB:AH): 1 + número de cortes ≤ cumplimiento; 0 si no hay dato
    cortes = [float(x) for x in p["cortes_cumplimiento"]]
    for c, s in zip(KPIS_CUMPL, KPIS_SC):
        df[s] = [0 if math.isnan(x) else 1 + sum(x >= q for q in cortes) for x in df[c]]

    # 6. Índice de Ejecución (AI:AL) con pesos calculados como en 'Análisis de KPIs'
    pesos = _pesos_ejecucion(df, ins.params_kpi)
    tiene = df[KPIS_CUMPL].notna().to_numpy()
    df["Cobertura de datos"] = tiene @ np.array(pesos)
    suma = df[KPIS_SC].to_numpy() @ np.array(pesos)
    activa_eleg = (df["Estatus"] == "Activa") & eleg
    df["Índice de Ejecución (0-100)"] = [
        (s / cob - 1) / 4 * 100 if (ok and cob >= p["min_cobertura_madurez"]) else math.nan
        for s, cob, ok in zip(suma, df["Cobertura de datos"], activa_eleg)]

    def clasif_madurez(est_, el, ie):
        if est_ != "Activa":
            return "Ruta inactiva"
        if el != "Sí":
            return "No elegible"
        if math.isnan(ie):
            return SIN_DATOS
        return "Madura" if ie >= p["corte_madura"] else ("Inmadura" if ie < p["corte_inmadura"] else "En desarrollo")
    df["Clasificación de madurez"] = [clasif_madurez(a, b, c) for a, b, c in
                                      zip(df["Estatus"], df["Elegible p/ estándar"], df["Índice de Ejecución (0-100)"])]
    df["Confiabilidad del dato"] = [
        "n/a" if math.isnan(ie) else ("Alta" if cob >= p["confiabilidad_alta"] else
                                      ("Media" if cob >= p["confiabilidad_media"] else "Baja"))
        for ie, cob in zip(df["Índice de Ejecución (0-100)"], df["Cobertura de datos"])]

    # 8. Saturación (AT:BE). Los deltas de la base son resta simple: celda vacía = 0
    df["Clientes actuales"] = r["clientes_act"].map(_num)
    df["Clientes potencial"] = r["clientes_pot"].map(_num)
    df["Delta clientes"] = [_n(a) - _n(b) for a, b in zip(r["clientes_pot"], r["clientes_act"])]
    df["Horas actuales"] = r["horas_act"].map(_num)
    df["Jornada completa"] = r["jornada"].map(_num)
    df["Delta horas"] = [_n(a) - _n(b) for a, b in zip(r["jornada"], r["horas_act"])]
    df["Capacidad instalada"] = r["cap_inst"].map(_num)
    df["Capacidad utilizada"] = r["cap_util"].map(_num)
    df["Delta capacidad"] = [_n(a) - _n(b) for a, b in zip(r["cap_inst"], r["cap_util"])]
    h = p["holgura_max"]

    def senales(ca, ci, dc, dh, dk):
        if math.isnan(ca) or ca <= 0 or _n(ci) <= 0:
            return math.nan
        return float((dc <= h) + (dh <= h) + (dk <= h))
    df["Señales sin holgura (0-3)"] = [senales(*t) for t in zip(df["Clientes actuales"], df["Capacidad instalada"],
                                                                df["Delta clientes"], df["Delta horas"],
                                                                df["Delta capacidad"])]
    df["Índice de Saturación"] = ["N/D" if math.isnan(x) else ("Saturada" if x >= 1 else "No saturada")
                                  for x in df["Señales sin holgura (0-3)"]]
    df["% Utilización de capacidad"] = [_div(u, i) for u, i in zip(df["Capacidad utilizada"], df["Capacidad instalada"])]

    # 9. Atractividad de la zona (BF:BQ) con el caso elegido
    c = ins.casos_atractividad[caso]
    disponible = ins.caso_disponible(caso)
    hog = r["hogares"].map(_num)
    if disponible:
        abc, pot = c["hog_abc"].map(_n), c["hog_pot"].map(_n)
        df[ZONA[0]] = [_div(a - b, hh) for a, b, hh in zip(abc, pot, hog)]
        df[ZONA[1]] = [_div(a, hh) for a, hh in zip(abc, hog)]
        df[ZONA[2]] = [_div(_n(i), b + hh) for i, b, hh in zip(c["pot_ingreso"], pot, hog)]
        df[ZONA[3]] = [_div(_n(i), b + hh) for i, b, hh in zip(c["pot_cu"], pot, hog)]
    else:
        for z in ZONA:
            df[z] = math.nan
    cnt = df[ZONA].notna().sum(axis=1)
    df["Fuente del dato de zona"] = np.where(cnt == 4, "Ruta (completo)", np.where(cnt == 0, "Sin dato", "Ruta (parcial)"))
    for z, s in zip(ZONA, ZONA_SC):
        df[s] = _puntaje_quintil_cedi(df, z)
    pa = np.array([float(x) for x in p["pesos_atractivo"]])
    df["Cobertura de datos de zona"] = df[ZONA].notna().to_numpy() @ pa
    suma_z = df[ZONA_SC].to_numpy() @ pa
    df["Índice de Atractivo (0-100)"] = [
        math.nan if cob < p["min_cobertura_zona"] else (s / cob - 1) / 4 * 100
        for s, cob in zip(suma_z, df["Cobertura de datos de zona"])]
    df["Nivel de atractivo"] = [
        SIN_DATOS if math.isnan(x) else ("Alto" if x >= p["corte_atractivo_alto"] else
                                         ("Bajo" if x < p["corte_atractivo_bajo"] else "Medio"))
        for x in df["Índice de Atractivo (0-100)"]]

    # 10. Resultados (BR:BW)
    df["Eje Y ruta - Madurez y saturación"] = [
        math.nan if (math.isnan(ie) or sat == "N/D") else
        ie * p["peso_ejecucion"] + (100 if sat == "Saturada" else 0) * p["peso_saturacion"]
        for ie, sat in zip(df["Índice de Ejecución (0-100)"], df["Índice de Saturación"])]
    df["Eje X ruta - Atractividad propia"] = df["Índice de Atractivo (0-100)"]
    corte_y = _redondear(float(df["Eje Y ruta - Madurez y saturación"].median()), 1)
    corte_x = _redondear(float(df["Eje X ruta - Atractividad propia"].median()), 1)

    def pq(y, x):
        if math.isnan(y) or math.isnan(x) or math.isnan(corte_x):
            return SIN_DATOS
        if y >= corte_y:
            return "P1" if x >= corte_x else "P2"
        return "P3" if x >= corte_x else "P4"
    df["P homologada de la ruta"] = [pq(y, x) for y, x in zip(df["Eje Y ruta - Madurez y saturación"],
                                                            df["Eje X ruta - Atractividad propia"])]
    df["Eje: estado de la ruta"] = [
        "No evaluable" if k in ("Ruta inactiva", "No elegible", SIN_DATOS) else ("Al tope" if s == "Saturada" else "Con espacio")
        for k, s in zip(df["Clasificación de madurez"], df["Índice de Saturación"])]
    df["Eje: potencial de la zona"] = [
        "No evaluable" if e == "No evaluable" else (
            "Sin datos de zona" if nv == SIN_DATOS else ("Zona atractiva" if nv == "Alto" else "Zona de bajo potencial"))
        for e, nv in zip(df["Eje: estado de la ruta"], df["Nivel de atractivo"])]
    nombres = {"P1": "P1 - Ampliar cobertura", "P2": "P2 - Mantener", "P3": "P3 - Desarrollar", "P4": "P4 - Evaluar"}
    df["Resultado"] = df["P homologada de la ruta"].map(nombres).fillna(SIN_DATOS)

    # Columnas adicionales para el árbol de decisión (no existen en la hoja de análisis)
    df["Hogares por manzana"] = r["hog_manzana"].map(_num)
    df["% A/B + C+"] = r["pct_abc_nse"].map(_num)
    df["Capacidad adicional (clientes)"] = r["cap_adicional"].map(_num)
    df["% Venta OP (ingreso)"] = r["pct_venta_op"].map(_num)
    # Hogares A/B + C+ del caso de potencial elegido (bloque "Atractividad Caso NN"), para el detalle del mapa
    df[COL_HOG_ABC_TOTAL] = c["hog_abc"].map(_num) if disponible else math.nan
    df[COL_HOG_ABC_POT] = c["hog_pot"].map(_num) if disponible else math.nan
    # Actuales = totales − potenciales (el bloque del caso trae totales y potenciales)
    df[COL_HOG_ABC_ACT] = (df[COL_HOG_ABC_TOTAL] - df[COL_HOG_ABC_POT]).clip(lower=0)
    df[COL_WEB] = r["pct_venta_web"].map(_num)
    # Para el árbol de decisión
    df[COL_ABC_CAPTURABLE] = [_div(pp, t) if t and t > 0 else math.nan
                              for pp, t in zip(df[COL_HOG_ABC_POT], df[COL_HOG_ABC_TOTAL])]
    df[COL_OP_VOLUMEN] = [_div(_num(o), _num(t)) for o, t in zip(r["cu_op"], r["cu_total"])]

    resumen = {
        "escenario": escenario, "caso": caso, "caso_disponible": disponible,
        "corte_madurez": corte_y, "corte_atractividad": corte_x,
        "pesos_ejecucion": dict(zip(KPIS, pesos)),
    }
    return df, resumen
